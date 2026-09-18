"""Evidence-grounded teaching reflection diagnosis and transformation."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
import json
import os
from pathlib import Path
import re
import time
from typing import Any, Iterable

from openai import OpenAI

from rag.bm25 import BM25Retriever, SearchHit
from rag.run_rag import MODEL, REASONING_EFFORT
from theory_subject_scope import (
    infer_subject_scope,
    subject_scope_label,
    theory_is_compatible,
)


ROOT = Path(__file__).resolve().parent
REFLECTION_ROOT = ROOT / "教学反思"
MAX_SOURCE_CHARS = 80_000
MAX_DIAGNOSES = 6
MAX_MODEL_SOURCE_CHARS = 18_000
MAX_OBSERVATION_CHARS = 6_000


def _sanitize_observation_value(value: Any, *, depth: int = 0) -> Any:
    """Keep observation summaries small and data-only before retrieval/model use."""

    if depth > 3:
        return "[内容已截断]"
    if isinstance(value, dict):
        cleaned: dict[str, Any] = {}
        for raw_key, raw_value in list(value.items())[:40]:
            key = str(raw_key).strip()[:80]
            if not key:
                continue
            item = _sanitize_observation_value(raw_value, depth=depth + 1)
            if item not in (None, "", [], {}):
                cleaned[key] = item
        return cleaned
    if isinstance(value, (list, tuple)):
        return [
            item
            for item in (_sanitize_observation_value(raw, depth=depth + 1) for raw in list(value)[:40])
            if item not in (None, "", [], {})
        ]
    if isinstance(value, (str, int, float, bool)):
        text = str(value).strip() if isinstance(value, str) else value
        return text[:1200] if isinstance(text, str) else text
    return str(value).strip()[:1200]


def _observation_context(payload: dict[str, Any]) -> dict[str, Any] | None:
    """Return an explicitly enabled classroom-observation supplement.

    Observation data is supplied by a completed classroom-analysis request.  It
    is deliberately kept separate from reflection facts: the model may use it
    to prioritize or explain a diagnosis, but must not turn it into invented
    events, counts, or learning outcomes.
    """

    if not bool(payload.get("include_observation")):
        return None
    raw = payload.get("observation_summary")
    if raw in (None, "", {}, []):
        return None
    cleaned = _sanitize_observation_value(raw)
    if cleaned in (None, "", {}, []):
        return None
    serialized = json.dumps(cleaned, ensure_ascii=False, separators=(",", ":"))
    if not serialized.strip() or serialized in {"null", "{}", "[]", '""'}:
        return None
    serialized = serialized[:MAX_OBSERVATION_CHARS]
    return {
        "data": cleaned,
        "serialized": serialized,
        "notice": "课堂观察摘要仅作为补充证据，不替代反思原文；摘要未提供的信息不得推断或虚构。",
    }


ISSUE_SPECS: tuple[dict[str, Any], ...] = (
    {
        "id": "pace",
        "category": "课堂节奏与目标达成",
        "pattern": r"时间|耗时|压缩|前松后紧|节奏|来不及|拖堂|仓促",
        "title": "核心学习任务与课堂时间配置需要进一步校准",
        "theory_hints": ("教学过程最优化理论", "认知负荷理论", "教学设计理论"),
        "interpretation": "反思文本显示教学活动之间存在时间竞争，需要区分核心学习任务与可压缩环节，并用学生学习证据决定时间配置。",
        "action": "为关键环节设置时间节点和停止条件，只保留能呈现不同思路的汇报，把节省的时间用于概念解释、迁移任务和即时反馈。",
        "talk": "先独立形成答案，小组只汇报不同的方法和理由；重复观点不再展开。",
        "indicator": "核心任务按计划完成；总结、迁移和反馈均有完整时间；重复汇报次数下降。",
    },
    {
        "id": "depth",
        "category": "认知深度与概念理解",
        "pattern": r"不深入|理解不够|浅|本质|概念|误区|追问|不会解释|只会|机械|照搬",
        "title": "学生表现尚不足以证明概念本质已经形成",
        "theory_hints": ("认知冲突教学理论", "概念转变理论", "脚手架教学理论", "深度学习理论"),
        "interpretation": "能完成操作或给出答案不等于形成稳定理解，需要用解释、反例、比较和迁移证据检验学生是否真正重构概念。",
        "action": "围绕原文暴露的典型认识设计反例和变式，连续追问依据、边界与条件变化，要求学生用课题中的具体内容完成解释。",
        "talk": "你的结论依据是什么？换一个条件还成立吗？请用一个反例或新例子检验。",
        "indicator": "多数学生能用学科证据解释概念、辨析反例，并在变式任务中说明方法为何成立。",
    },
    {
        "id": "participation",
        "category": "学生参与与学习机会",
        "pattern": r"参与|少数|个别|基础薄弱|学困|内向|展示|表达能力强|举手|沉默|两极",
        "title": "课堂表达与学习机会在学生之间分布不均",
        "theory_hints": ("差异化教学理论", "合作学习理论", "教育公平理论", "最近发展区理论"),
        "interpretation": "统一任务和自愿举手容易使可见表现集中于少数学生，需要同时调整任务难度、表达支架和发言机制。",
        "action": "采用先写后说、分层问题、随机抽取与角色轮换，为不同准备度学生配置可完成且必须留下证据的任务。",
        "talk": "每个人先写下一条依据，再由组内不同角色依次说明、补充和质疑。",
        "indicator": "每名学生至少形成一项可检查产出；发言覆盖率提高；同一学生不连续承担核心汇报。",
    },
    {
        "id": "collaboration",
        "category": "合作学习质量",
        "pattern": r"小组|合作|讨论|同伴|汇报|分工|交流|倾听",
        "title": "小组活动需要从形式参与转向共同认知产出",
        "theory_hints": ("合作学习理论", "社会互动学习理论", "群体动力理论"),
        "interpretation": "出现小组活动并不能证明合作学习发生，需检查是否有积极互赖、个体责任、观点比较与共同成果。",
        "action": "明确记录员、解释员和质疑员等角色，规定共同产出，并要求汇报小组如何处理不同意见。",
        "talk": "请说明你们有哪些不同看法，最后依据什么形成共同结论。",
        "indicator": "每组提交共同成果和个人痕迹；汇报包含观点比较；角色在不同轮次轮换。",
    },
    {
        "id": "authenticity",
        "category": "情境真实性与知识迁移",
        "pattern": r"生活|真实|应用|举例|迁移|实践|情境|脱离|联系实际|决策",
        "title": "生活联结尚需转化为学生真实解决问题的过程",
        "theory_hints": ("情境认知理论", "学习迁移理论", "项目式学习理论", "RME现实数学教育"),
        "interpretation": "教师举例或展示生活素材主要建立表面联系，真正迁移需要学生在新情境中选择知识、解释依据并评价结果。",
        "action": "将课题内容嵌入信息完整、目标明确的真实任务，让学生经历提取证据、形成方案、作出判断和反思边界。",
        "talk": "在这个真实任务中，你为什么选择本课的方法？它能解决什么，不能说明什么？",
        "indicator": "学生能独立识别适用条件、用知识支持判断，并说明结论在新情境中的边界。",
    },
    {
        "id": "feedback",
        "category": "评价反馈与证据闭环",
        "pattern": r"评价|反馈|证据|效果|达成|检测|观察|判断|作业|练习|掌握",
        "title": "改进设想缺少可观察、可复核的效果证据",
        "theory_hints": ("形成性评价理论", "反馈理论", "教学评价理论"),
        "interpretation": "经验判断只有转化为具体学习证据和达成标准，才能在下一轮课堂中判断策略是否有效。",
        "action": "为每项改进配置学生行为、作品或互动数据，明确收集时点、达成标准和未达标时的调整方式。",
        "talk": "我们用哪一项作品或课堂表现判断这次调整是否真正帮助了学习？",
        "indicator": "每项行动至少对应一条学生证据、一个量化或等级标准和一项未达标调整。",
    },
    {
        "id": "teacher_dominance",
        "category": "师生关系与学习主体",
        "pattern": r"讲得多|教师讲|灌输|包办|替代|牵着走|被动|主体|自主|生成",
        "title": "教师支持与学生自主建构之间需要重新平衡",
        "theory_hints": ("建构主义学习理论", "主体性教育理论", "发现学习理论"),
        "interpretation": "教师提供过多路径或结论可能降低任务的认知责任，学生需要经历提出想法、比较证据和修正观点的过程。",
        "action": "把结论前移为学生任务，教师用问题、材料和反馈提供支架，在学生形成可比较观点后再组织归纳。",
        "talk": "先不要猜老师的答案，请用自己的证据提出一种解释，再比较同伴方案。",
        "indicator": "学生提出的方法和问题数量增加；教师连续讲解时长下降；学生能说明观点修正过程。",
    },
    {
        "id": "goals",
        "category": "教学目标与任务一致性",
        "pattern": r"目标|重难点|偏离|任务|活动|达成|核心素养|知识点|教学内容",
        "title": "目标、学习任务与评价证据的一致性需要加强",
        "theory_hints": ("教学目标分类理论", "逆向教学设计理论", "核心素养理论"),
        "interpretation": "反思中的目标描述需要落实为学生可完成的任务和可观察的表现，避免活动丰富但无法证明目标达成。",
        "action": "把每个核心目标改写为学生表现，逐项检查任务是否提供练习机会、评价是否采集对应证据。",
        "talk": "完成这项任务后，你能用什么表现证明自己真正学会了？",
        "indicator": "核心目标均对应学习任务和评价证据；无关活动被删除；学生能复述成功标准。",
    },
)


OUTPUT_TYPES = {
    "reflection": "教学反思深化版",
    "case": "理论化教学案例",
    "research": "一课一研教研报告",
    "paper": "教学研究论文框架",
    "lesson": "课例研究方案",
    "experiment": "下一轮教学改进方案",
}


@dataclass(frozen=True)
class ReflectionReferences:
    design: str
    original_sample: str
    revised_sample: str
    cases: str


def _docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        document = Document(str(path))
    except (OSError, ValueError):
        return ""
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def extract_reflection_document(filename: str, data: bytes) -> dict[str, str]:
    """Extract editable text from common reflection document formats."""

    name = str(filename or "reflection.txt")
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        text = ""
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text:
            text = data.decode("utf-8", errors="replace")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("服务端缺少 python-docx，无法解析 Word 文件") from exc
        try:
            document = Document(BytesIO(data))
        except (OSError, ValueError) as exc:
            raise ValueError("Word 文件损坏或不是有效的 DOCX") from exc
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text = "\n".join(blocks)
    elif suffix == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("服务端缺少 PyMuPDF，无法解析 PDF 文件") from exc
        try:
            with fitz.open(stream=data, filetype="pdf") as document:
                text = "\n".join(page.get_text("text") for page in document)
        except (RuntimeError, ValueError) as exc:
            raise ValueError("PDF 文件损坏、加密或无法提取正文") from exc
    else:
        raise ValueError("仅支持 DOCX、PDF、TXT、MD 和 CSV 文件")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        raise ValueError("文件中没有提取到可用正文")
    metadata = infer_reflection_metadata(text, name)
    return {"filename": name, "text": text[:MAX_SOURCE_CHARS], **metadata}


def load_reflection_references(root: Path = REFLECTION_ROOT) -> ReflectionReferences:
    design_parts = [_docx_text(path) for path in (root / "教学反思设计方案").glob("*.docx")]
    sample_root = root / "教学反思生成样例"
    original = next(sample_root.glob("*平均数*教学反思.docx"), None)
    revised = next(sample_root.glob("基于落地理论修改后*.docx"), None)
    case_parts = [_docx_text(path) for path in (root / "教学反思案例").glob("*.docx")]
    return ReflectionReferences(
        design="\n".join(part for part in design_parts if part)[:18_000],
        original_sample=_docx_text(original)[:12_000] if original else "",
        revised_sample=_docx_text(revised)[:14_000] if revised else "",
        cases="\n".join(part for part in case_parts if part)[:14_000],
    )


def infer_reflection_metadata(source: str, filename: str = "") -> dict[str, str]:
    source = str(source or "")
    clean_filename = re.sub(r"\.(?:docx|pdf|txt|md|csv)$", "", str(filename or ""), flags=re.I)
    clean_filename = re.sub(r"(?:教学反思|课后反思|听评课记录)$", "", clean_filename).strip(" _-—")
    candidates = [clean_filename]
    candidates.extend(re.findall(r"(?:执教|教学|学习|上完|讲授)?\s*((?:人教版|统编版|北师大版|苏教版)?[^。\n]{0,24}《[^》]{1,30}》)", source[:2500]))
    candidates.extend(f"《{item}》" for item in re.findall(r"《([^》]{1,30})》", source[:2500]))
    lesson = next((re.sub(r"^(?:执教|教学|学习|上完|讲授)\s*", "", item).strip() for item in candidates if item and len(item) <= 80), "")
    if not lesson:
        heading = next((line.strip() for line in source.splitlines() if 2 <= len(line.strip()) <= 60), "")
        lesson = re.sub(r"(?:教学反思|课后反思)$", "", heading).strip() or "未命名课例"
    topic_match = re.search(r"《([^》]+)》", lesson)
    topic = topic_match.group(1).strip() if topic_match else re.sub(r"(?:教学反思|课后反思)", "", lesson).strip()
    if not topic:
        topic = "本课"
    project = f"{topic}课例持续改进"
    return {"lesson_name": lesson, "project_name": project}


def _sentences(source: str) -> list[str]:
    return [item.strip() for item in re.split(r"(?<=[。！？；])|\n+", source) if len(item.strip()) >= 6]


def _lesson_topic(lesson_name: str, source: str = "") -> str:
    match = re.search(r"《([^》]{1,40})》", str(lesson_name or ""))
    if not match:
        match = re.search(r"《([^》]{1,40})》", str(source or "")[:3000])
    return match.group(1).strip() if match else ""


def _quoted_topics(text: str) -> set[str]:
    return {item.strip() for item in re.findall(r"《([^》]{1,40})》", str(text or "")) if item.strip()}


def _has_unrelated_lesson(text: str, current_topic: str) -> bool:
    if not current_topic:
        return False
    return any(topic != current_topic for topic in _quoted_topics(text))


def _theory_excerpt(hit: SearchHit, fallback: str, current_topic: str) -> str:
    """Return theory mechanism text without borrowing another lesson's case."""

    chunk = hit.chunk
    text = str(chunk.get("text", "")).strip()
    source_type = str(chunk.get("source_type", ""))
    if source_type == "课堂范本" and _has_unrelated_lesson(text, current_topic):
        return fallback

    useful: list[str] = []
    labels = ("理论定位", "内涵", "核心观点", "理论说明", "关键判断标准", "一句话记忆")
    for line in text.splitlines():
        normalized = re.sub(r"\s+", " ", line).strip()
        if not normalized or _has_unrelated_lesson(normalized, current_topic):
            continue
        if normalized.startswith(labels):
            useful.append(normalized)
    if useful:
        excerpt = "\n".join(useful)
    else:
        generic = re.split(
            r"(?:课堂范本|案例\s*[一二三四五六七八九十\d]*\s*[:：]|课堂落地方式\s*[:：]|例如\s*[,，]?\s*在?《)",
            text,
            maxsplit=1,
        )[0]
        lines = [
            re.sub(r"\s+", " ", line).strip()
            for line in generic.splitlines()
            if line.strip() and not _has_unrelated_lesson(line, current_topic)
        ]
        excerpt = "\n".join(lines)
    # A header-only classroom chunk does not explain a mechanism.
    if len(re.sub(r"理论名称|分类路径|来源类型|原始标题|正文", "", excerpt)) < 80:
        return fallback
    return excerpt[:560]


def _evidence_for(source: str, spec: dict[str, Any]) -> tuple[str, bool]:
    pattern = re.compile(str(spec["pattern"]))
    matches = [sentence for sentence in _sentences(source) if pattern.search(sentence)]
    if matches:
        return "；".join(matches[:2])[:500], True
    return "反思原文未提供可直接支持该判断的课堂事实，需补充学生表现、作品或过程数据。", False


def _theory_hits(
    query: str,
    retriever: BM25Retriever,
    theory_names: Iterable[str] = (),
    top_k: int = 8,
    current_topic: str = "",
    subject_scope: str = "",
) -> list[SearchHit]:
    requested = {str(name).strip() for name in theory_names if str(name).strip()}
    hits = retriever.search(query, top_k=max(top_k * 3, 18))
    filtered: list[SearchHit] = []
    seen: set[str] = set()
    for hit in hits:
        name = str(hit.chunk.get("theory_name", "")).strip()
        if not name or name in seen:
            continue
        if not theory_is_compatible(hit.chunk, subject_scope):
            continue
        hit_text = str(hit.chunk.get("text", ""))
        if (
            current_topic
            and str(hit.chunk.get("source_type", "")) == "课堂范本"
            and _has_unrelated_lesson(hit_text, current_topic)
        ):
            continue
        if requested and name not in requested:
            continue
        coverage = retriever.query_coverage(query, hit)
        if hit.score < 24.0 and coverage < 0.1 and name not in requested:
            continue
        seen.add(name)
        filtered.append(hit)
        if len(filtered) >= top_k:
            break
    return filtered


def _source_item(hit: SearchHit) -> dict[str, Any]:
    chunk = hit.chunk
    return {
        "chunk_id": str(chunk.get("chunk_id", "")),
        "theory_name": str(chunk.get("theory_name", "")),
        "category_path": list(chunk.get("category_path", [])),
        "source_type": str(chunk.get("source_type", "")),
        "source_title": str(chunk.get("source_title", "")),
        "score": round(float(hit.score), 4),
    }


def _safe_json(text: str) -> dict[str, Any] | None:
    value = str(text or "").strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*|\s*```$", "", value, flags=re.I | re.S).strip()
    start, end = value.find("{"), value.rfind("}")
    if start < 0 or end <= start:
        return None
    try:
        parsed = json.loads(value[start : end + 1])
    except json.JSONDecodeError:
        return None
    return parsed if isinstance(parsed, dict) else None


def _model_timeout(stage: str, default: float = 60.0) -> float:
    name = f"REFLECTION_{stage.upper()}_TIMEOUT_SECONDS"
    try:
        return max(15.0, float(os.getenv(name, str(default))))
    except ValueError:
        return default


def _structured_model_json(
    client: Any,
    *,
    instructions: str,
    prompt: str,
    max_output_tokens: int,
    timeout: float,
) -> tuple[dict[str, Any] | None, Any | None]:
    """Request JSON and retry only an unusable model payload at low effort."""

    configured_effort = os.getenv("REFLECTION_REASONING_EFFORT", REASONING_EFFORT).strip().lower()
    if configured_effort not in {"minimal", "low", "medium", "high", "xhigh"}:
        configured_effort = REASONING_EFFORT
    efforts = (configured_effort, "low") if configured_effort != "low" else ("low",)
    model_client = client.with_options(timeout=timeout, max_retries=0) if hasattr(client, "with_options") else client
    last_response: Any | None = None
    for effort in efforts:
        arguments: dict[str, Any] = {
            "model": MODEL,
            "instructions": instructions,
            "input": prompt,
            "reasoning": {"effort": effort},
            "text": {"format": {"type": "json_object"}},
            "max_output_tokens": max_output_tokens,
            "store": False,
        }
        try:
            last_response = model_client.responses.create(**arguments)
        except TypeError:
            # Preserve compatibility with small test doubles and older SDKs.
            arguments.pop("text", None)
            last_response = model_client.responses.create(**arguments)
        parsed = _safe_json(getattr(last_response, "output_text", ""))
        if parsed is not None:
            return parsed, last_response
    return None, last_response


def _fallback_diagnosis(payload: dict[str, Any], retriever: BM25Retriever) -> dict[str, Any]:
    source = str(payload.get("source_text", ""))[:MAX_SOURCE_CHARS].strip()
    if len(source) < 30:
        raise ValueError("请先输入至少 30 字的教学反思原文")
    observation = _observation_context(payload)
    # Observation metrics can help identify which issue specs to inspect, but
    # they remain separate from reflection facts and evidence-status labels.
    signal_source = source
    if observation:
        signal_source = f"{source}\n{observation['serialized']}"
    metadata = infer_reflection_metadata(source, str(payload.get("filename", "")))
    if str(payload.get("lesson_name", "")).strip():
        metadata["lesson_name"] = str(payload["lesson_name"]).strip()
    if str(payload.get("project_name", "")).strip():
        metadata["project_name"] = str(payload["project_name"]).strip()
    current_topic = _lesson_topic(metadata["lesson_name"], source)
    subject_scope = infer_subject_scope(
        context="\n".join([
            metadata["lesson_name"],
            str(payload.get("filename", "")),
            source,
        ])
    )
    subject_label = subject_scope_label(subject_scope)
    matched = [spec for spec in ISSUE_SPECS if re.search(spec["pattern"], signal_source)]
    if len(matched) < 3:
        matched.extend(spec for spec in ISSUE_SPECS if spec not in matched)
    selected_specs = matched[:MAX_DIAGNOSES]
    diagnoses: list[dict[str, Any]] = []
    recommendation_by_name: dict[str, dict[str, Any]] = {}
    references: dict[str, dict[str, Any]] = {}
    for index, spec in enumerate(selected_specs, start=1):
        evidence, supported = _evidence_for(source, spec)
        query = " ".join([
            subject_label, metadata["lesson_name"], spec["category"], spec["title"], evidence,
            *spec["theory_hints"], "课堂观察指标 改进策略 AI理论识别规则 关键判断标准",
        ])
        if observation:
            # Include the user-selected observation summary in retrieval so
            # relevant theory chunks can be ranked, while keeping it labelled
            # as supplementary evidence.
            query = f"{query} 课堂观察补充证据 {observation['serialized'][:1800]}"
        hits = _theory_hits(
            query,
            retriever,
            top_k=3,
            current_topic=current_topic,
            subject_scope=subject_scope,
        )
        primary = hits[0] if hits else None
        theory_name = str(primary.chunk.get("theory_name", "")) if primary else str(spec["theory_hints"][0])
        mechanism = _theory_excerpt(primary, str(spec["interpretation"]), current_topic) if primary else str(spec["interpretation"])
        source_items = [_source_item(hit) for hit in hits]
        for item in source_items:
            references[item["chunk_id"] or f"{item['source_title']}:{len(references)}"] = item
        confidence = max(58, min(96, int(70 + (primary.score if primary else 0) / 4)))
        diagnosis = {
            "id": str(spec["id"]),
            "rank": index,
            "category": str(spec["category"]),
            "title": str(spec["title"]),
            "priority": "高" if index <= 3 else "中",
            "evidence": evidence,
            "professional_judgment": str(spec["interpretation"]),
            "status": "原文有证据" if supported else "证据待补",
            "theory": theory_name,
            "theory_explanation": mechanism,
            "action": str(spec["action"]),
            "talk": str(spec["talk"]),
            "indicator": str(spec["indicator"]),
            "confidence": confidence,
            "sources": source_items,
        }
        diagnoses.append(diagnosis)
        for hit in hits:
            name = str(hit.chunk.get("theory_name", "")).strip()
            if not name:
                continue
            existing = recommendation_by_name.get(name)
            score = max(55, min(99, int(68 + hit.score / 3)))
            if existing and existing["score"] >= score:
                continue
            recommendation_by_name[name] = {
                "name": name,
                "group": " → ".join(hit.chunk.get("category_path", [])) or "教育理论库",
                "score": score,
                "core": _theory_excerpt(hit, str(spec["interpretation"]), current_topic)[:420],
                "mechanism": str(spec["interpretation"]),
                "action": str(spec["action"]),
                "selected": index <= 4,
                "sources": [_source_item(hit)],
            }
        if theory_name not in recommendation_by_name:
            recommendation_by_name[theory_name] = {
                "name": theory_name,
                "group": "教育理论库",
                "score": confidence,
                "core": mechanism,
                "mechanism": str(spec["interpretation"]),
                "action": str(spec["action"]),
                "selected": index <= 4,
                "sources": source_items,
            }
    recommended = sorted(recommendation_by_name.values(), key=lambda item: item["score"], reverse=True)[:8]
    selected_names = [item["name"] for item in recommended[:4]]
    for item in recommended:
        item["selected"] = item["name"] in selected_names
    strengths = []
    strength_rules = (
        (r"情境|真实|生活", "关注情境与经验联结"),
        (r"探究|操作|讨论|合作", "能复盘学生学习过程"),
        (r"发现|不足|问题|改进", "具有问题意识和改进意愿"),
        (r"目标|重难点|核心素养", "能够回看目标达成"),
    )
    for pattern, label in strength_rules:
        if re.search(pattern, source):
            strengths.append(label)
    return {
        **metadata,
        "summary": (
            f"从反思原文识别 {len(diagnoses)} 个核心问题，并从理论库召回 {len(recommended)} 个候选理论。"
            + ("已纳入课堂观察摘要作为补充证据。" if observation else "")
        ),
        "strengths": strengths or ["能够记录教学过程并提出改进方向"],
        "diagnoses": diagnoses,
        "recommended_theories": recommended,
        "selected_theories": selected_names,
        "references": list(references.values()),
        "used_model": False,
        "warning": None,
        "model": MODEL,
        "observation_included": bool(observation),
        "observation_summary": observation["data"] if observation else None,
        "observation_notice": observation["notice"] if observation else "未合并课堂观察数据。",
    }


def _diagnosis_prompt(payload: dict[str, Any], fallback: dict[str, Any], refs: ReflectionReferences) -> str:
    compact_diagnoses = [
        {
            field: item.get(field)
            for field in (
                "id",
                "category",
                "title",
                "priority",
                "evidence",
                "professional_judgment",
                "status",
                "theory",
                "theory_explanation",
                "action",
                "talk",
                "indicator",
                "confidence",
            )
        }
        for item in fallback.get("diagnoses", [])
    ]
    compact_input = {
        "source_text": str(payload.get("source_text", ""))[:MAX_MODEL_SOURCE_CHARS],
        "project_name": fallback.get("project_name", ""),
        "lesson_name": fallback.get("lesson_name", ""),
        "source_type": payload.get("source_type", ""),
        "teacher_focus": str(payload.get("teacher_focus", ""))[:2400],
        "strengths": fallback.get("strengths", []),
        "diagnoses": compact_diagnoses,
    }
    observation = _observation_context(payload)
    observation_block = ""
    if observation:
        compact_input["observation_supplement"] = observation["data"]
        observation_block = f"""
课堂观察补充证据（用户已完成的课堂观察摘要，仅作辅助，不替代反思原文）：
{observation["serialized"]}
使用边界：只能把上述摘要作为补充线索来排序或解释问题；摘要未提供的课堂事实、学生人数、频次、作品和效果不得推断或虚构。
"""
    return f"""你是严谨的教师专业反思诊断专家。请只返回一个合法 JSON 对象，不输出 Markdown 或思考过程。

任务：依据教师反思原文和已检索理论证据，精修专业判断与改进行动。不得编造课堂事件、学生人数、频次、作品或实施效果。课堂事实、证据状态、理论名称和来源由系统锁定，你不得修改。

JSON 顶层只输出：summary, strengths, diagnoses。
diagnoses 每项只输出：id, category, title, priority(高/中/低), professional_judgment, theory_explanation, action, talk, indicator, confidence(0-100)。
要求：保持候选 id 和顺序；核心问题具体指向本课；判断不能超出给定事实；行动必须能在下一节课直接执行；indicator 必须可观察或可计数。不要输出 recommended_theories、selected_theories、sources、evidence、status 或 theory。

当前课例数据：{json.dumps(compact_input, ensure_ascii=False)}
{observation_block}
设计要求摘要：{refs.design[:1800]}
"""


def _merge_diagnosis(model_data: dict[str, Any], fallback: dict[str, Any]) -> dict[str, Any]:
    result = dict(fallback)
    for field in ("summary",):
        value = model_data.get(field)
        if isinstance(value, str) and value.strip():
            result[field] = value.strip()
    if isinstance(model_data.get("strengths"), list):
        strengths = [str(item).strip() for item in model_data["strengths"] if str(item).strip()]
        if strengths:
            result["strengths"] = strengths[:6]
    if result.get("observation_included") and "补充证据" not in str(result.get("summary", "")):
        result["summary"] = f"{str(result.get('summary', '')).rstrip()} 已纳入课堂观察摘要作为补充证据。".strip()
    model_items = [item for item in model_data.get("diagnoses", []) if isinstance(item, dict)]
    model_by_id = {
        str(item.get("id", "")): item
        for item in model_items
        if str(item.get("id", ""))
    }
    merged_diagnoses: list[dict[str, Any]] = []
    for index, base in enumerate(fallback["diagnoses"], start=1):
        raw = model_by_id.get(str(base.get("id", "")))
        if raw is None and index <= len(model_items) and not model_items[index - 1].get("id"):
            raw = model_items[index - 1]
        item = dict(base)
        if raw is not None:
            for field in ("category", "title", "professional_judgment", "theory_explanation", "action", "talk", "indicator"):
                value = raw.get(field)
                if isinstance(value, str) and value.strip():
                    item[field] = value.strip()
            priority = str(raw.get("priority", item["priority"]))
            item["priority"] = priority if priority in {"高", "中", "低"} else item["priority"]
            try:
                item["confidence"] = max(0, min(100, int(raw.get("confidence", item["confidence"]))))
            except (TypeError, ValueError):
                pass
        item["rank"] = index
        # Facts, evidence boundaries, theory selection and provenance remain deterministic.
        item["evidence"] = base["evidence"]
        item["status"] = base["status"]
        item["theory"] = base["theory"]
        item["sources"] = base["sources"]
        merged_diagnoses.append(item)
    result["diagnoses"] = merged_diagnoses[:MAX_DIAGNOSES]
    # Theory recommendations and default selection are retrieval results, not model output.
    result["recommended_theories"] = fallback["recommended_theories"]
    result["selected_theories"] = fallback["selected_theories"]
    return result


def diagnose_reflection(
    payload: dict[str, Any],
    retriever: BM25Retriever,
    client: OpenAI | None = None,
    references: ReflectionReferences | None = None,
) -> dict[str, Any]:
    fallback = _fallback_diagnosis(payload, retriever)
    if client is None or not bool(payload.get("use_model", True)):
        return fallback
    refs = references or load_reflection_references()
    try:
        parsed, _response = _structured_model_json(
            client,
            instructions="你是教师专业反思诊断专家。严格输出合法 JSON，不展示内部思维链。",
            prompt=_diagnosis_prompt(payload, fallback, refs),
            max_output_tokens=5600,
            timeout=_model_timeout("diagnose"),
        )
        if not parsed:
            fallback["warning"] = "模型两次返回均无法解析，已保留理论库证据诊断。"
            return fallback
        current_topic = _lesson_topic(str(fallback.get("lesson_name", "")), str(payload.get("source_text", "")))
        if _has_unrelated_lesson(json.dumps(parsed, ensure_ascii=False), current_topic):
            fallback["warning"] = "模型诊断混入了其他课例，已保留当前课例的理论库证据诊断。"
            return fallback
        result = _merge_diagnosis(parsed, fallback)
        result["used_model"] = True
        return result
    except Exception as exc:
        fallback["warning"] = f"模型诊断暂不可用，已返回理论库证据诊断：{type(exc).__name__}"
        return fallback


def _fallback_actions(diagnosis: dict[str, Any], selected_theories: list[str]) -> dict[str, Any]:
    selected = set(selected_theories)
    theory_profiles = {
        str(item.get("name", "")).strip(): item
        for item in diagnosis.get("recommended_theories", [])
        if isinstance(item, dict) and str(item.get("name", "")).strip()
    }
    actions: list[dict[str, Any]] = []
    for index, item in enumerate(diagnosis.get("diagnoses", []), start=1):
        original_theory = str(item.get("theory", ""))
        theory = original_theory
        if selected and theory not in selected:
            alternative = selected_theories[(index - 1) % len(selected_theories)]
            theory = alternative or theory
        profile = theory_profiles.get(theory, {})
        action = str(item.get("action", ""))
        theory_action = str(profile.get("action", "")).strip()
        if theory_action and theory != original_theory:
            action = f"针对“{item.get('title', '课堂问题')}”，依据{theory}落实以下调整：{theory_action}"
        actions.append({
            "id": item.get("id", f"action-{index}"),
            "index": index,
            "title": action.split("，", 1)[0] or "下一轮教学调整",
            "issue": str(item.get("title", "课堂问题")),
            "theory": theory,
            "action": action,
            "talk": str(item.get("talk", "")),
            "indicator": str(item.get("indicator", "")),
            "success_criteria": str(item.get("indicator", "")),
            "evidence_to_collect": "记录相关学生话语、学习单或课堂作品，并与本轮事实证据对照。",
        })
    goal = "；".join(item["issue"] for item in actions[:3])
    strategy = "\n".join(f"{item['index']}. {item['action']}" for item in actions)
    next_round = {
        "goal": f"下一轮重点改善：{goal}" if goal else "围绕核心课堂问题开展下一轮实践验证。",
        "strategy": strategy,
        "indicators": [item["indicator"] for item in actions],
    }
    return {"actions": actions, "next_round": next_round, "used_model": False, "warning": None, "model": MODEL}


def generate_reflection_actions(payload: dict[str, Any], client: OpenAI | None = None) -> dict[str, Any]:
    diagnosis = dict(payload.get("diagnosis") or {})
    selected = [str(item).strip() for item in payload.get("selected_theories", []) if str(item).strip()]
    fallback = _fallback_actions(diagnosis, selected)
    if client is None or not bool(payload.get("use_model", True)):
        return fallback
    compact_diagnosis = {
        "project_name": diagnosis.get("project_name", ""),
        "lesson_name": diagnosis.get("lesson_name", ""),
        "diagnoses": [
            {
                field: item.get(field)
                for field in (
                    "id", "category", "title", "priority", "evidence", "status",
                    "professional_judgment", "theory", "theory_explanation",
                )
            }
            for item in diagnosis.get("diagnoses", [])
        ],
        "selected_theory_profiles": [
            {
                field: item.get(field)
                for field in ("name", "core", "mechanism", "action")
            }
            for item in diagnosis.get("recommended_theories", [])
            if item.get("name") in selected
        ],
    }
    prompt = f"""只返回合法 JSON。根据已确认理论组合，把诊断转化为下一节课可执行的行动。
顶层只输出 actions。actions 每项只输出 id,theory,action,talk,indicator,evidence_to_collect。
约束：行动必须具体到本课任务、提问、组织或评价；不得编造原文事实；理论必须来自已确认列表；每项行动都要有可观察指标和证据采集方式。
已确认理论：{json.dumps(selected, ensure_ascii=False)}
当前课例诊断：{json.dumps(compact_diagnosis, ensure_ascii=False)}
"""
    try:
        parsed, _response = _structured_model_json(
            client,
            instructions="你是教学改进行动设计专家。严格输出合法 JSON。",
            prompt=prompt,
            max_output_tokens=3600,
            timeout=_model_timeout("actions", 45.0),
        )
        if not parsed or not isinstance(parsed.get("actions"), list):
            fallback["warning"] = "模型两次返回的行动方案格式均无效，已返回证据化行动草案。"
            return fallback
        current_topic = _lesson_topic(
            str(diagnosis.get("lesson_name", "")),
            str(payload.get("source_text", "")),
        )
        if _has_unrelated_lesson(json.dumps(parsed, ensure_ascii=False), current_topic):
            fallback["warning"] = "模型行动方案混入了其他课例，已返回当前课例的行动草案。"
            return fallback
        raw_actions = [item for item in parsed["actions"][:8] if isinstance(item, dict)]
        raw_by_id = {
            str(item.get("id", "")): item
            for item in raw_actions
            if str(item.get("id", ""))
        }
        actions = []
        for index, base in enumerate(fallback["actions"], start=1):
            raw = raw_by_id.get(str(base.get("id", "")))
            if raw is None and index <= len(raw_actions) and not raw_actions[index - 1].get("id"):
                raw = raw_actions[index - 1]
            item = dict(base)
            if raw is not None:
                for field in ("title", "theory", "action", "talk", "indicator", "success_criteria", "evidence_to_collect"):
                    if isinstance(raw.get(field), str) and raw[field].strip():
                        item[field] = raw[field].strip()
            if selected and item.get("theory") not in selected:
                item["theory"] = selected[0]
            item["id"] = base.get("id", f"action-{index}")
            item["issue"] = base.get("issue", "课堂问题")
            item["index"] = index
            actions.append(item)
        if not actions:
            return fallback
        result = dict(fallback)
        result["actions"] = actions
        result["next_round"] = {
            "goal": f"下一轮重点改善：{'；'.join(item['issue'] for item in actions[:3])}",
            "strategy": "\n".join(f"{item['index']}. {item['action']}" for item in actions),
            "indicators": [item["indicator"] for item in actions],
        }
        result["used_model"] = True
        return result
    except Exception as exc:
        fallback["warning"] = f"模型行动设计暂不可用，已返回本地草案：{type(exc).__name__}"
        return fallback


def _fallback_outcome(payload: dict[str, Any]) -> str:
    diagnosis = dict(payload.get("diagnosis") or {})
    actions_data = dict(payload.get("actions_data") or {})
    source = str(payload.get("source_text", "")).strip()
    project = str(payload.get("project_name") or diagnosis.get("project_name") or "课堂持续改进")
    lesson = str(payload.get("lesson_name") or diagnosis.get("lesson_name") or "本课")
    output_type = str(payload.get("output_type", "reflection"))
    selected = [str(item) for item in payload.get("selected_theories", [])]
    diagnoses = list(diagnosis.get("diagnoses", []))
    actions = list(actions_data.get("actions", []))
    title = OUTPUT_TYPES.get(output_type, OUTPUT_TYPES["reflection"])
    source_excerpt = source[:6000]
    if len(source) > len(source_excerpt):
        source_excerpt += "\n\n（原文较长，此处节选；完整原文保留在材料区。）"
    evidence = "\n".join(f"{index}. {item.get('evidence', '')}\n   专业判断：{item.get('professional_judgment', item.get('title', ''))}" for index, item in enumerate(diagnoses, start=1))
    profiles = {
        str(item.get("name", "")): item
        for item in diagnosis.get("recommended_theories", [])
        if isinstance(item, dict) and item.get("name")
    }
    theory_items = selected or [str(item.get("theory", "")) for item in diagnoses if item.get("theory")]
    theory_text_parts: list[str] = []
    for name in dict.fromkeys(theory_items):
        profile = profiles.get(name, {})
        explanation = str(profile.get("mechanism") or profile.get("core") or "").strip()
        if not explanation:
            explanation = next(
                (str(item.get("theory_explanation", "")) for item in diagnoses if item.get("theory") == name),
                "",
            )
        theory_text_parts.append(f"{len(theory_text_parts) + 1}. {name}：{explanation}")
    theory_text = "\n".join(theory_text_parts)
    action_text = "\n".join(f"{index}. {item.get('action', '')}\n   教师话术：{item.get('talk', '')}\n   观察指标：{item.get('indicator', '')}" for index, item in enumerate(actions, start=1))
    indicators = "\n".join(f"- {item.get('indicator', '')}" for item in actions) or "- 为每项行动补充学生表现、作品或互动数据。"
    focus = str(payload.get("teacher_focus", "")).strip() or "围绕本轮高优先级问题验证教学调整是否改善学生学习。"

    if output_type == "case":
        return f"""# {lesson}{title}

## 一、教学背景与问题情境

本案例属于“{project}”，以教师真实反思为第一手材料，聚焦：{focus}

## 二、课堂实施与事实证据

{evidence or '原文证据不足，需补充课堂事实。'}

## 三、核心问题

{chr(10).join(f'{index}. {item.get("title", "")}' for index, item in enumerate(diagnoses, start=1))}

## 四、理论依据与机制分析

{theory_text or '尚未确认理论组合。'}

## 五、改进后的课堂实施

{action_text or '尚未生成下一轮行动。'}

## 六、学生学习证据与效果判断

{indicators}

## 七、案例启示

案例结论仅由已记录事实支持。下一轮需保留学生话语、作品和任务完成数据，检验理论所解释的机制是否真正发生。
"""

    if output_type == "research":
        return f"""# {lesson}{title}

项目：{project}

## 一、核心教研问题

{focus}

## 二、课堂事实与问题基线

{evidence or '原文证据不足，需补充课堂事实。'}

## 三、理论解释

{theory_text or '尚未确认理论组合。'}

## 四、教研改进共识

{action_text or '尚未生成下一轮行动。'}

## 五、下一轮观察分工与指标

{indicators}

## 六、教研结论与后续安排

下一轮按相同指标采集证据，分别判断已改善、持续存在、新出现和暂无证据的问题，再决定保留、调整或替换策略。
"""

    if output_type == "paper":
        return f"""# {title}：基于课堂证据与理论介入的{lesson}持续改进研究

## 一、问题提出

研究聚焦：{focus}

课堂问题基线：
{evidence or '原文证据不足，需补充课堂事实。'}

## 二、理论框架

{theory_text or '尚未确认理论组合。'}

## 三、研究问题

1. 当前课堂问题通过何种教学与学习机制产生？
2. 基于选定理论的行动如何改变学生学习过程？
3. 哪些课堂证据能够支持或否定改进效果？

## 四、研究设计

采用课例研究与行动研究相结合的路径，以教师反思、课堂话语、学生作品和任务完成情况为证据，开展“诊断—理论介入—再次实践—指标比较”的多轮研究。

## 五、实践干预与分析维度

{action_text or '尚未生成实践干预。'}

## 六、预期证据与判断标准

{indicators}

## 七、证据边界

当前材料只能形成研究假设，不能预先宣称策略有效；结论须经下一轮课堂数据验证。
"""

    if output_type == "lesson":
        return f"""# {lesson}{title}

## 一、课例研究主题

{focus}

## 二、第一轮问题基线

{evidence or '原文证据不足，需补充课堂事实。'}

## 三、理论介入

{theory_text or '尚未确认理论组合。'}

## 四、教学改进设计

{action_text or '尚未生成下一轮行动。'}

## 五、再次实践与观察分工

{indicators}

## 六、跨轮比较规则

逐项记录原问题是否改善、是否持续存在、是否出现新问题，以及哪些判断仍缺少证据。不同轮次使用一致指标，避免只凭总体感受比较。

## 七、成果沉淀

保存每轮反思、理论选择、教学方案、课堂证据和判断结论，逐步形成可复核的课例研究档案。
"""

    if output_type == "experiment":
        return f"""# {lesson}{title}

## 一、改进目标

{actions_data.get('next_round', {}).get('goal') or focus}

## 二、理论依据

{theory_text or '尚未确认理论组合。'}

## 三、问题基线

{evidence or '原文证据不足，需补充课堂事实。'}

## 四、实施措施与教师话术

{action_text or '尚未生成下一轮行动。'}

## 五、课堂观察指标

{indicators}

## 六、证据采集与效果判断

每项行动对应收集学生话语、学习单、课堂作品或互动数据。达到标准则标记“已改善”，连续未达标则标记“持续存在”，材料不足则标记“暂无证据”，不得以教师主观感受代替结果。
"""

    return f"""# {lesson}{title}

项目：{project}

## 一、原始反思与课堂现象

{source_excerpt}

## 二、课堂事实与核心问题

{evidence or '原文证据不足，需补充课堂事实。'}

## 三、理论解释

{theory_text or '尚未确认理论组合。'}

## 四、改进目标与实施策略

{action_text or '尚未生成下一轮行动。'}

## 五、下一轮观察与效果判断

{indicators}

## 六、专业反思结论

本次转化不以理论替代教师经验，而是用课堂事实界定问题、用理论解释机制、用可观察指标检验改进。下一轮实践应保留学生话语、作品和任务完成情况，区分已改善、持续存在、新出现与证据不足的问题。
"""


def _ensure_outcome_heading(text: str, lesson_name: str, output_type: str) -> str:
    """Lock the editable outcome heading to the current lesson and type."""

    body = str(text or "").strip()
    lesson = str(lesson_name or "本课").strip() or "本课"
    title = OUTPUT_TYPES.get(output_type, OUTPUT_TYPES["reflection"])
    desired = f"# {lesson}{title}"
    lines = body.splitlines()
    first_content = next((index for index, line in enumerate(lines) if line.strip()), None)
    if first_content is not None and re.match(r"^#\s+", lines[first_content].strip()):
        if lesson in lines[first_content] and title in lines[first_content]:
            return body
        lines[first_content] = desired
        return "\n".join(lines).strip()
    return f"{desired}\n\n{body}" if body else desired


def generate_reflection_outcome(
    payload: dict[str, Any],
    client: OpenAI | None = None,
    references: ReflectionReferences | None = None,
) -> dict[str, Any]:
    fallback = _fallback_outcome(payload)
    if client is None or not bool(payload.get("use_model", True)):
        return {"content": fallback, "used_model": False, "warning": None, "model": MODEL}
    refs = references or load_reflection_references()
    output_type = str(payload.get("output_type", "reflection"))
    diagnosis = dict(payload.get("diagnosis") or {})
    actions_data = dict(payload.get("actions_data") or {})
    compact_payload = {
        "source_text": str(payload.get("source_text", ""))[:MAX_MODEL_SOURCE_CHARS],
        "project_name": payload.get("project_name") or diagnosis.get("project_name", ""),
        "lesson_name": payload.get("lesson_name") or diagnosis.get("lesson_name", ""),
        "output_type": output_type,
        "writing_style": payload.get("writing_style", "teacher"),
        "teacher_focus": str(payload.get("teacher_focus", ""))[:2400],
        "selected_theories": payload.get("selected_theories", []),
        "diagnoses": [
            {
                field: item.get(field)
                for field in (
                    "category", "title", "priority", "evidence", "status",
                    "professional_judgment", "theory", "theory_explanation",
                )
            }
            for item in diagnosis.get("diagnoses", [])
        ],
        "actions": [
            {
                field: item.get(field)
                for field in (
                    "issue", "theory", "action", "talk", "indicator",
                    "success_criteria", "evidence_to_collect",
                )
            }
            for item in actions_data.get("actions", [])
        ],
        "next_round": actions_data.get("next_round", {}),
    }
    prompt = f"""请生成“{OUTPUT_TYPES.get(output_type, OUTPUT_TYPES['reflection'])}”，只输出 Markdown 正文，不输出 JSON 或思考过程。
必须围绕当前课例和原始反思中的具体内容展开，不得套用《平均数》事实；理论只使用已确认组合，相关判断必须有原文证据。结构应覆盖课堂现象、事实证据、核心问题、理论如何解释、可直接执行的改进、教师话术、观察指标和证据边界。详细、专业、可继续编辑。

当前课例数据：{json.dumps(compact_payload, ensure_ascii=False)}
设计要求摘要：{refs.design[:1800]}
"""
    try:
        configured_effort = os.getenv("REFLECTION_REASONING_EFFORT", REASONING_EFFORT).strip().lower()
        if configured_effort not in {"minimal", "low", "medium", "high", "xhigh"}:
            configured_effort = REASONING_EFFORT
        efforts = (configured_effort, "low") if configured_effort != "low" else ("low",)
        timeout = _model_timeout("outcome", 80.0)
        model_client = client.with_options(timeout=timeout, max_retries=0) if hasattr(client, "with_options") else client
        current_topic = _lesson_topic(
            str(payload.get("lesson_name") or (payload.get("diagnosis") or {}).get("lesson_name", "")),
            str(payload.get("source_text", "")),
        )
        last_issue = "incomplete"
        for effort in efforts:
            response = model_client.responses.create(
                model=MODEL,
                instructions="你是教师专业成果编辑。基于证据与理论写作，不虚构课堂事实。",
                input=prompt,
                reasoning={"effort": effort},
                max_output_tokens=7600,
                store=False,
            )
            text = str(getattr(response, "output_text", "") or "").strip()
            if _has_unrelated_lesson(text, current_topic):
                last_issue = "unrelated"
                continue
            if len(text) < 800 or "理论" not in text or "改进" not in text:
                last_issue = "incomplete"
                continue
            text = _ensure_outcome_heading(
                text,
                str(payload.get("lesson_name") or diagnosis.get("lesson_name") or "本课"),
                output_type,
            )
            return {"content": text, "used_model": True, "warning": None, "model": MODEL}
        warning = (
            "模型两次生成均混入其他课例，已返回当前课例的结构化草案。"
            if last_issue == "unrelated"
            else "模型两次生成的成果内容均不完整，已返回结构化草案。"
        )
        return {"content": fallback, "used_model": False, "warning": warning, "model": MODEL}
    except Exception as exc:
        return {"content": fallback, "used_model": False, "warning": f"模型成果生成暂不可用：{type(exc).__name__}", "model": MODEL}


def build_growth_profile(payload: dict[str, Any]) -> dict[str, Any]:
    rounds = [item for item in payload.get("rounds", []) if isinstance(item, dict)]
    diagnosis = dict(payload.get("diagnosis") or {})
    actions = list((payload.get("actions_data") or {}).get("actions", []))
    diagnoses = list(diagnosis.get("diagnoses", []))
    categories = Counter(str(item.get("category", "课堂诊断")) for item in diagnoses)
    theories = Counter(str(item.get("theory", "")) for item in diagnoses if item.get("theory"))
    evidence_ready = sum(str(item.get("status", "")).find("证据待补") < 0 for item in diagnoses)
    dimensions = [
        {"name": "课堂诊断", "score": min(100, 58 + len(diagnoses) * 7)},
        {"name": "理论解释", "score": min(100, 55 + len(theories) * 9)},
        {"name": "行动设计", "score": min(100, 55 + len(actions) * 8)},
        {"name": "证据评价", "score": min(100, 50 + evidence_ready * 8)},
        {"name": "持续研究", "score": min(100, 48 + len(rounds) * 10)},
        {"name": "成果表达", "score": 72 if payload.get("outcome_content") else 56},
    ]
    persistent = [item.get("title", "") for item in diagnoses if item.get("priority") == "高"][:3]
    return {
        "dimensions": dimensions,
        "insights": [
            {"label": "高频关注领域", "value": "、".join(name for name, _ in categories.most_common(3)) or "待积累更多反思"},
            {"label": "长期改善方向", "value": "从经验描述走向课堂事实、理论机制与指标验证"},
            {"label": "持续挑战问题", "value": "、".join(persistent) or "尚无跨轮证据"},
            {"label": "理论使用特点", "value": "、".join(name for name, _ in theories.most_common(3)) or "尚未确认理论"},
            {"label": "下一成长主题", "value": "在下一轮课堂收集学生学习证据并验证行动效果"},
        ],
        "rounds": rounds,
        "generated_at": time.strftime("%Y-%m-%d"),
    }
